import os
import math
from datetime import datetime, timedelta
from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
import requests
import openpyxl

app = Flask(__name__)
CORS(app)

# ---------------- CONFIGURATION ----------------
TELEGRAM_BOT_TOKEN = "8730808658:AAE2CxCZ6m2dQqqFxu7RuFgyMSyonE2NNCk"  
TELEGRAM_CHAT_ID = "-1003954771903"            

COMPANY_LAT = 39.6548
COMPANY_LON = 66.9597
MAX_ALLOWED_DISTANCE_METERS = 150

EXCEL_FILE = "davomat.xlsx"

# OYLIK QOIDALARI:
BASE_MONTHLY_SALARY = 4000000  # 26 kunlik norma uchun 4 mln so'm
STANDARD_WORKING_DAYS = 26
DAILY_RATE = BASE_MONTHLY_SALARY / STANDARD_WORKING_DAYS  # Ortiqcha kunlik haq

# Excel faylni yaratish
if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Davomat"
    ws.append([
        "Ism Familiya", "Smena", "Kelgan Vaqti", "Ketgan Vaqti", 
        "Jami Soat", "Kechikish (daq)", "Vaqtli Ketish (daq)", 
        "Hisoblangan Oylik (so'm)", "Masofa (m)"
    ])
    wb.save(EXCEL_FILE)

# ---------------- WEB ROUTELAR ----------------
@app.route('/')
def serve_index():
    return send_from_directory('.', 'index.html')

@app.route('/admin')
def serve_admin():
    return send_from_directory('.', 'admin.html')

@app.route('/<path:filename>')
def serve_static_files(filename):
    return send_from_directory('.', filename)

# ---------------- ADMIN LOGIN & DATA API ----------------
@app.route('/admin-login', methods=['POST'])
def admin_login():
    data = request.json
    if data.get('username') == 'admin' and data.get('password') == 'elsevar2026':
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 401

@app.route('/get-davomat', methods=['GET'])
def get_davomat():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Davomat"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(list(row))
    return jsonify(rows)

@app.route('/download-excel', methods=['GET'])
def download_excel():
    if os.path.exists(EXCEL_FILE):
        return send_file(EXCEL_FILE, as_attachment=True)
    return "Excel fayl topilmadi!", 404

# ---------------- GEOLOKATSIYA HISOoblash ----------------
def calculate_distance(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi/2)**2 + math.cos(phi1)*math.cos(phi2) * math.sin(dlambda/2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c

# ---------------- 3 TA SMENANI ANIQLASH ----------------
def get_shift_info(dt):
    hour = dt.hour
    today = dt.date()
    if 7 <= hour < 15:
        return "1-Smena (08:00-16:00)", datetime.combine(today, datetime.strptime("08:00", "%H:%M").time()), datetime.combine(today, datetime.strptime("16:00", "%H:%M").time())
    elif 15 <= hour < 23:
        return "2-Smena (16:00-00:00)", datetime.combine(today, datetime.strptime("16:00", "%H:%M").time()), datetime.combine(today, datetime.strptime("00:00", "%H:%M").time()) + timedelta(days=1)
    else:
        if hour == 23:
            return "3-Smena (00:00-08:00)", datetime.combine(today + timedelta(days=1), datetime.strptime("00:00", "%H:%M").time()), datetime.combine(today + timedelta(days=1), datetime.strptime("08:00", "%H:%M").time())
        else:
            return "3-Smena (00:00-08:00)", datetime.combine(today, datetime.strptime("00:00", "%H:%M").time()), datetime.combine(today, datetime.strptime("08:00", "%H:%M").time())

# ---------------- FACE-ID & DAVOMAT API ----------------
@app.route('/send-face', methods=['POST'])
def send_face():
    try:
        full_name = request.form.get('full_name')
        user_lat = float(request.form.get('latitude', 0))
        user_lon = float(request.form.get('longitude', 0))
        action_type_btn = request.form.get('action_type', 'in') # 'in' yoki 'out'
        photo = request.files.get('photo')

        distance = calculate_distance(COMPANY_LAT, COMPANY_LON, user_lat, user_lon)
        if distance > MAX_ALLOWED_DISTANCE_METERS:
            return jsonify({
                "status": "error", 
                "message": f"Siz korxonadan {int(distance)}m uzoqdasiz! Face-ID faqat korxona hududida ishlaydi."
            }), 400

        now = datetime.now()
        now_time_str = now.strftime("%Y-%m-%d %H:%M:%S")
        time_display = now.strftime("%H:%M:%S")
        shift_name, shift_start, shift_end = get_shift_info(now)

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["Davomat"]

        action_type = "Ishga keldi"
        worked_hours = 0
        salary = BASE_MONTHLY_SALARY
        late_minutes = 0
        early_leave_minutes = 0

        if action_type_btn == 'in':
            if now > shift_start:
                late_minutes = int((now - shift_start).total_seconds() // 60)

            ws.append([
                full_name, shift_name, now_time_str, None, 
                0, late_minutes, 0, BASE_MONTHLY_SALARY, round(distance, 1)
            ])
            action_type = "Ishga keldi"
        else:
            action_type = "Ishdan ketdi"
            row_found = False
            for row in range(ws.max_row, 1, -1):
                emp_name = ws.cell(row=row, column=1).value
                time_in_str = ws.cell(row=row, column=3).value
                time_out_val = ws.cell(row=row, column=4).value

                if emp_name == full_name and time_out_val is None and time_in_str:
                    time_in = datetime.strptime(time_in_str, "%Y-%m-%d %H:%M:%S")
                    ws.cell(row=row, column=4, value=now_time_str)
                    
                    duration_seconds = (now - time_in).total_seconds()
                    worked_hours = round(duration_seconds / 3600.0, 2)

                    if now < shift_end:
                        early_leave_minutes = int((shift_end - now).total_seconds() // 60)

                    user_total_days = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == full_name)
                    
                    if user_total_days > STANDARD_WORKING_DAYS:
                        extra_days = user_total_days - STANDARD_WORKING_DAYS
                        salary = round(BASE_MONTHLY_SALARY + (extra_days * DAILY_RATE))
                    else:
                        salary = BASE_MONTHLY_SALARY

                    ws.cell(row=row, column=5, value=worked_hours)
                    ws.cell(row=row, column=7, value=early_leave_minutes)
                    ws.cell(row=row, column=8, value=salary)
                    
                    shift_name = ws.cell(row=row, column=2).value
                    row_found = True
                    break
            
            if not row_found:
                return jsonify({"status": "error", "message": "Avval 'Ishga kelish' tugmasini bosib ro'yxatdan o'ting!"}), 400

        wb.save(EXCEL_FILE)

        # Telegramga yuborish
        caption_text = (
            f"👤 <b>Xodim:</b> {full_name}\n"
            f"🔄 <b>Smena:</b> {shift_name}\n"
            f"📌 <b>Holat:</b> {action_type}\n"
            f"⏰ <b>Vaqt:</b> {time_display}\n"
            f"📍 <b>Masofa:</b> {int(distance)} metr\n"
        )
        if action_type == "Ishga keldi" and late_minutes > 0:
            caption_text += f"⚠️ <b>Kechikish:</b> {late_minutes} daqiqa\n"
        elif action_type == "Ishdan ketdi":
            if early_leave_minutes > 0:
                caption_text += f"⚠️ <b>Vaqtli ketish:</b> {early_leave_minutes} daqiqa oldin\n"
            caption_text += f"⏱ <b>Jami ishladi:</b> {worked_hours} soat\n💰 <b>Jami Oylik:</b> {salary:,} so'm\n"

        if TELEGRAM_BOT_TOKEN != "YANGI_BOT_TOKENINGIZNI_SHU_YERGA_YOZING" and photo:
            requests.post(
                f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendPhoto",
                data={"chat_id": TELEGRAM_CHAT_ID, "caption": caption_text, "parse_mode": "HTML"},
                files={"photo": (photo.filename, photo.read(), photo.content_type)}
            )

        return jsonify({"status": "success", "action": action_type, "message": "Muvaffaqiyatli saqlandi!"})

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)